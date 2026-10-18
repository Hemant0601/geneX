#!/usr/bin/env python3
"""Checks lead-leak-calculator.xlsx by recomputing every result independently.

For each scenario it (1) writes the inputs into a temporary copy of the workbook, (2) lets a real
spreadsheet engine recalculate it (LibreOffice headless, if installed; and the Python `formulas`
package, if installed) and (3) compares every Results cell with a hand-written Python model of
the same maths. Exit code 1 on any mismatch.

Run:  python3 lead-magnet/src/verify_calculator.py
"""
import math
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.dont_write_bytecode = True  # keep __pycache__ out of the repo
sys.path.insert(0, str(HERE))
import build_calculator as bc  # noqa: E402

XLSX = HERE.parent / "lead-leak-calculator.xlsx"


def model(i):
    """Same maths as the sheet, written out by hand. i: dict input-number -> value (None = blank)."""
    n, ah, ahl, dl, nf, vr, ns, cr, dv, hrs, wb, cost = (i[k] for k in range(1, 13))
    r = {}
    r["n"] = n
    r["ah"] = n * ah
    r["ahl"] = n * ah * ahl
    r["dl"] = n * (1 - ah) * dl
    r["late"] = r["ahl"] + r["dl"]
    r["latepct"] = r["late"] / n if n else 0
    r["nf"] = n * nf
    p_late, p_nf = r["latepct"], nf
    r["risk"] = n * (1 - (1 - p_late) * (1 - p_nf))       # P(late or no follow-up), independent
    r["ns"] = n * vr * ns
    r["hrs"] = r["ns"] * hrs
    r["deals"] = n * cr
    r["rev"] = r["deals"] * dv
    r["rdeals"] = r["risk"] * cr
    r["rrev"] = r["rdeals"] * dv
    r["rrevy"] = r["rrev"] * 12
    if wb is None:
        r["wdeals"], r["wrev"], r["wrevy"] = "Enter input 11 first", "—", "—"
    else:
        r["wdeals"] = r["rdeals"] * wb
        r["wrev"] = r["wdeals"] * dv
        r["wrevy"] = r["wrev"] * 12
    r["be"] = "—" if cost is None or dv == 0 else cost / dv
    r["bem"] = "—" if (cost is None or wb in (None, 0) or cr == 0 or dv == 0 or n == 0) else cost / r["wrev"]
    return r


EXAMPLE = {k + 1: v[1] for k, v in enumerate(bc.INPUTS[:12])}
SCENARIOS = {
    "A: example numbers, no guess, no cost": dict(EXAMPLE),
    "B: 60 enquiries, 40% after hours, guess 10%, fix ₹34,999": {
        1: 60, 2: 0.40, 3: 0.90, 4: 0.25, 5: 0.60, 6: 0.12, 7: 0.25, 8: 0.03, 9: 120000, 10: 1.5,
        11: 0.10, 12: 34999,
    },
    "C: edge case, zero enquiries": {**EXAMPLE, 1: 0, 11: 0.2, 12: 9999},
}


def write_inputs(src, dst, inputs):
    wb = load_workbook(src)
    ws = wb["Your numbers"]
    for k, v in inputs.items():
        ws.cell(row=bc.R[k], column=3, value=v)
    wb.save(dst)


def recalc_libreoffice(path, outdir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    profile = Path(outdir) / "lo-profile"
    cmd = [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(Path(outdir) / "lo"), str(path)]
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    out = Path(outdir) / "lo" / Path(path).name
    wb = load_workbook(out, data_only=True)
    rs = wb["Results"]
    return {k: rs.cell(row=row, column=3).value for k, row in bc.RESULT_ROWS.items()}, rs["B3"].value


def recalc_formulas_pkg(path):
    try:
        import formulas  # noqa: F401
    except ImportError:
        return None
    import contextlib
    import io
    import formulas
    with contextlib.redirect_stderr(io.StringIO()):
        xl = formulas.ExcelModel().loads(str(path)).finish()
        sol = {str(k).upper(): v for k, v in xl.calculate().items()}
    fname = Path(path).name.upper()
    out = {}
    for k, row in bc.RESULT_ROWS.items():
        v = sol.get(f"'[{fname}]RESULTS'!C{row}")
        v = v.value[0, 0] if v is not None and hasattr(v, "value") else v
        out[k] = v.item() if hasattr(v, "item") else v
    return out


def same(a, b):
    if isinstance(b, str) or isinstance(a, str):
        return str(a) == str(b)
    if a is None or b is None:
        return False
    return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=1e-9)


def main():
    bc.build()  # rebuild so RESULT_ROWS is populated and the file is current
    failures = 0
    with tempfile.TemporaryDirectory() as tmp:
        for name, inputs in SCENARIOS.items():
            full = {**{k: None for k in range(1, 13)}, **inputs}
            path = Path(tmp) / f"scenario-{name[0]}.xlsx"
            write_inputs(XLSX, path, full)
            expect = model(full)
            engines = {}
            lo = recalc_libreoffice(path, tmp)
            if lo:
                engines["LibreOffice"], banner = lo
            fp = recalc_formulas_pkg(path)
            if fp:
                engines["formulas"] = fp
            print(f"\n{name}")
            if not engines:
                print("  (no spreadsheet engine available; only the Python model ran)")
            for key in bc.RESULT_ROWS:
                got = {e: vals.get(key) for e, vals in engines.items()}
                ok = all(same(v, expect[key]) for v in got.values())
                failures += 0 if ok else 1
                ev = expect[key]
                shown = f"{ev:,.2f}" if isinstance(ev, (int, float)) else ev
                flag = "ok " if ok else "BAD"
                extra = "" if ok else f"   engines: {got}"
                print(f"  {flag} {key:8s} {shown}{extra}")
            if lo:
                print(f"  banner: {banner}")
    print("\nAll results match." if not failures else f"\n{failures} mismatch(es).")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
